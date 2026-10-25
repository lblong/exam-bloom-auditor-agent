from docx import Document

cau_hoi = [
    # (thân câu hỏi, {A,B,C,D}, đáp án)
    ("Thủ đô của Việt Nam là gì?",
     {"A": "Hà Nội", "B": "Đà Nẵng", "C": "Huế", "D": "TP.HCM"}, "A"),
    ("Giải thích vì sao HTTP là giao thức phi trạng thái (stateless)?",
     {"A": "Vì server không lưu thông tin phiên giữa các request",
      "B": "Vì HTTP chạy trên UDP",
      "C": "Vì HTTP không có header",
      "D": "Vì HTTP chỉ dùng được một lần"}, "A"),
    ("Một trang web bị chậm khi có nhiều người truy cập cùng lúc. Bạn đề xuất giải pháp nào phù hợp nhất?",
     {"A": "Thêm cache và cân bằng tải",
      "B": "Xóa toàn bộ dữ liệu",
      "C": "Tắt server",
      "D": "Đổi tên miền"}, "A"),
    # LỖI 1: trùng phương án B và C
    ("Ngôn ngữ nào dùng để truy vấn cơ sở dữ liệu quan hệ?",
     {"A": "Java", "B": "SQL", "C": "SQL", "D": "HTML"}, "B"),
    # LỖI 2: vi phạm quy chế "Tất cả các câu trên đều đúng"
    ("Python là ngôn ngữ nào sau đây?",
     {"A": "Ngôn ngữ thông dịch", "B": "Ngôn ngữ đa năng",
      "C": "Có thư viện phong phú", "D": "Tất cả các câu trên đều đúng"}, "D"),
    # LỖI 3: thiếu đáp án đúng
    ("Cổng mặc định của HTTPS là bao nhiêu?",
     {"A": "80", "B": "443", "C": "21", "D": "25"}, None),
]

doc = Document()
for i, (than, pa, da) in enumerate(cau_hoi, start=1):
    doc.add_paragraph(f"Câu {i}. {than}")
    for k in ["A", "B", "C", "D"]:
        doc.add_paragraph(f"{k}. {pa[k]}")
    if da:
        doc.add_paragraph(f"Đáp án: {da}")
    # nếu da là None thì cố ý không ghi dòng Đáp án

doc.save("data/de_mau.docx")
print("Đã tạo data/de_mau.docx")