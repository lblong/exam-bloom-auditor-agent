import json
import random

from docx import Document

# ===== 30 CÂU CHUẨN: 7 NHO, 8 HIEU, 8 VAN_DUNG, 7 PHAN_TICH =====
CHUAN = [
    # ---- NHO (7) ----
    ("Viết tắt HTTP là gì?",
     ["Hypertext Transfer Protocol", "High Transfer Text Protocol", "Hyper Terminal Transfer Program", "Host Transfer Text Protocol"], "A", "NHO"),
    ("Cấu trúc dữ liệu nào hoạt động theo nguyên tắc LIFO?",
     ["Queue", "Stack", "Array", "Linked List"], "B", "NHO"),
    ("Từ khóa nào trong Java dùng để kế thừa một lớp?",
     ["implements", "inherits", "extends", "super"], "C", "NHO"),
    ("Câu lệnh SQL nào dùng để truy xuất dữ liệu từ bảng?",
     ["INSERT", "SELECT", "UPDATE", "DELETE"], "B", "NHO"),
    ("Cổng mặc định của giao thức HTTPS là bao nhiêu?",
     ["80", "21", "25", "443"], "D", "NHO"),
    ("Độ phức tạp thời gian của tìm kiếm nhị phân trên mảng đã sắp xếp là gì?",
     ["O(n)", "O(log n)", "O(n^2)", "O(1)"], "B", "NHO"),
    ("Lệnh Git nào dùng để tạo một commit mới?",
     ["git push", "git pull", "git commit", "git clone"], "C", "NHO"),

    # ---- HIEU (8) ----
    ("Giải thích vì sao HTTP được gọi là giao thức phi trạng thái (stateless)?",
     ["Vì server không lưu thông tin phiên giữa các request", "Vì HTTP chạy trên UDP", "Vì HTTP không có header", "Vì HTTP chỉ dùng được một lần"], "A", "HIEU"),
    ("Sự khác nhau cơ bản giữa Stack và Queue là gì?",
     ["Stack lấy phần tử vào trước ra trước, Queue lấy vào sau ra trước", "Stack lấy phần tử vào sau ra trước, Queue lấy vào trước ra trước", "Stack chỉ lưu số, Queue chỉ lưu chuỗi", "Stack nhanh hơn Queue ở mọi thao tác"], "B", "HIEU"),
    ("Tại sao nên dùng PreparedStatement thay vì Statement khi truy vấn SQL từ Java?",
     ["Vì Statement không chạy được trên MySQL", "Vì PreparedStatement luôn nhanh gấp 10 lần", "Vì PreparedStatement tham số hóa truy vấn nên ngăn được SQL Injection", "Vì PreparedStatement không cần kết nối CSDL"], "C", "HIEU"),
    ("Tính đóng gói (encapsulation) trong lập trình hướng đối tượng có ý nghĩa gì?",
     ["Cho phép một lớp có nhiều lớp cha", "Che giấu dữ liệu bên trong và chỉ cho truy cập qua các phương thức công khai", "Tự động giải phóng bộ nhớ của đối tượng", "Cho phép chạy nhiều luồng cùng lúc"], "B", "HIEU"),
    ("Khóa ngoại (foreign key) có vai trò gì trong cơ sở dữ liệu quan hệ?",
     ["Tăng tốc độ ghi dữ liệu", "Mã hóa dữ liệu trong bảng", "Thiết lập liên kết và ràng buộc toàn vẹn giữa hai bảng", "Tự động sao lưu dữ liệu"], "C", "HIEU"),
    ("Giải thích vì sao tìm kiếm nhị phân yêu cầu mảng phải được sắp xếp trước?",
     ["Vì mỗi bước loại bỏ một nửa dựa trên so sánh với phần tử giữa, điều này chỉ đúng khi mảng có thứ tự", "Vì mảng chưa sắp xếp thì không thể truy cập theo chỉ số", "Vì nhị phân chỉ làm việc với số chẵn", "Vì Java bắt buộc phải sắp xếp trước khi tìm"], "A", "HIEU"),
    ("Điểm khác nhau chính giữa abstract class và interface trong Java là gì?",
     ["Interface có thể chứa trường dữ liệu thay đổi được còn abstract class thì không", "Abstract class không thể có phương thức", "Abstract class có thể chứa trạng thái và phương thức đã cài đặt, còn interface chủ yếu định nghĩa hợp đồng cần tuân theo", "Hai khái niệm này hoàn toàn giống nhau"], "C", "HIEU"),
    ("Hàm băm (hash function) trong bảng băm có vai trò gì?",
     ["Sắp xếp các khóa theo thứ tự tăng dần", "Ánh xạ một khóa thành chỉ số vị trí lưu trong bảng", "Mã hóa khóa để bảo mật", "Xóa các khóa bị trùng"], "B", "HIEU"),

    # ---- VAN_DUNG (8) ----
    ("Bảng SinhVien có 10 dòng, trong đó 7 sinh viên có Diem >= 5. Câu lệnh SELECT COUNT(*) FROM SinhVien WHERE Diem >= 5 trả về giá trị nào?",
     ["3", "10", "5", "7"], "D", "VAN_DUNG"),
    ("Bạn cần viết chương trình kiểm tra chuỗi dấu ngoặc '(()())' có cân bằng hay không. Cấu trúc dữ liệu nào phù hợp nhất để cài đặt?",
     ["Stack", "Queue", "Binary Search Tree", "Hash Table"], "A", "VAN_DUNG"),
    ("Website của bạn bị chậm khi hàng nghìn người truy cập cùng lúc. Bạn chọn giải pháp nào phù hợp nhất?",
     ["Xóa toàn bộ dữ liệu", "Thêm cache và cân bằng tải", "Đổi tên miền", "Tắt server vào ban ngày"], "B", "VAN_DUNG"),
    ("Bạn cần lưu danh sách sinh viên và thường xuyên tìm theo mã số với thời gian gần như O(1). Nên dùng cấu trúc nào trong Java?",
     ["ArrayList", "LinkedList", "HashMap", "Stack"], "C", "VAN_DUNG"),
    ("Bạn lỡ sửa sai một file nhưng chưa commit và muốn hủy thay đổi trên file đó. Bạn dùng lệnh nào?",
     ["git push --force", "git restore tên_file", "git clone", "git merge"], "B", "VAN_DUNG"),
    ("Mỗi khách hàng có nhiều đơn hàng. Khi thiết kế bảng DonHang trong CSDL, bạn xử lý mối quan hệ này như thế nào?",
     ["Thêm cột khóa ngoại MaKhachHang vào bảng DonHang", "Gộp toàn bộ đơn hàng vào một cột của bảng KhachHang", "Xóa bảng KhachHang", "Dùng một bảng duy nhất không có khóa"], "A", "VAN_DUNG"),
    ("Áp dụng phân tích độ phức tạp cho đoạn mã gồm hai vòng for lồng nhau, mỗi vòng chạy từ 1 đến n. Độ phức tạp thời gian của nó là gì?",
     ["O(n)", "O(log n)", "O(n^2)", "O(2^n)"], "C", "VAN_DUNG"),
    ("Trong Spring Boot, bạn muốn API tại đường dẫn /students trả về dữ liệu JSON. Annotation nào phù hợp để đặt trên lớp controller?",
     ["@Entity", "@RestController", "@Repository", "@Configuration"], "B", "VAN_DUNG"),

    # ---- PHAN_TICH (7) ----
    ("Đoạn mã Java: String s = null; if (s.length() > 0) {...} gây ra lỗi khi chạy. Phân tích nguyên nhân gốc của lỗi này.",
     ["Biến s tham chiếu tới null nên gọi phương thức length() gây NullPointerException", "Java không cho khai báo kiểu String", "Phương thức length() không tồn tại", "Điều kiện > 0 luôn sai về cú pháp"], "A", "PHAN_TICH"),
    ("Hệ thống thường xuyên chèn và xóa phần tử ở đầu danh sách. Hãy phân tích vì sao danh sách liên kết hiệu quả hơn mảng trong trường hợp này.",
     ["Vì danh sách liên kết dùng ít bộ nhớ hơn trong mọi trường hợp", "Vì mảng không cho phép xóa phần tử", "Vì mảng phải dịch chuyển toàn bộ phần tử còn lại, còn danh sách liên kết chỉ cập nhật con trỏ", "Vì danh sách liên kết truy cập theo chỉ số nhanh hơn"], "C", "PHAN_TICH"),
    ("Truy vấn SELECT lọc theo cột email trên bảng 10 triệu dòng chạy rất chậm. Hãy phân tích nguyên nhân khả dĩ nhất và hướng xử lý.",
     ["Do tên cột quá dài, cần đổi tên cột", "Do chưa có chỉ mục (index) trên cột email nên phải quét toàn bộ bảng, cần tạo index", "Do bảng có quá nhiều cột, cần xóa bớt dữ liệu", "Do SQL không hỗ trợ lọc theo chuỗi"], "B", "PHAN_TICH"),
    ("So sánh Bubble sort O(n^2) và Merge sort O(n log n) khi sắp xếp 1 triệu phần tử. Phân tích lựa chọn nào hợp lý hơn và đánh đổi là gì.",
     ["Bubble sort vì không tốn thêm bộ nhớ nên luôn tốt hơn", "Hai thuật toán có hiệu năng như nhau với mọi n", "Merge sort vì số phép so sánh tăng chậm hơn nhiều, đổi lại cần thêm bộ nhớ phụ", "Bubble sort vì dễ viết hơn nên nhanh hơn"], "C", "PHAN_TICH"),
    ("Hàm đệ quy fib(n) = fib(n-1) + fib(n-2) chạy rất chậm với n = 45. Hãy phân tích nguyên nhân.",
     ["Do máy tính không đủ bộ nhớ cho biến n", "Do các bài toán con bị tính lặp đi lặp lại rất nhiều lần nên số lần gọi tăng theo hàm mũ", "Do Java không hỗ trợ đệ quy", "Do phép cộng trong Java chậm"], "B", "PHAN_TICH"),
    ("Hai thành viên cùng sửa cùng một dòng của một file trên hai nhánh Git khác nhau, khi merge phát sinh xung đột. Hãy phân tích vì sao Git không tự xử lý được.",
     ["Vì Git không hỗ trợ merge nhiều nhánh", "Vì file bị hỏng khi sửa", "Vì Git không có lịch sử commit", "Vì cùng một vùng nội dung bị thay đổi khác nhau ở hai nhánh nên Git không biết giữ phiên bản nào"], "D", "PHAN_TICH"),
    ("Trong bảng băm dùng linear probing, hiệu năng giảm mạnh khi hệ số tải (load factor) tiến gần 1. Hãy phân tích nguyên nhân.",
     ["Do các phần tử bị dồn thành cụm liên tiếp (clustering) nên số lần dò tìm tăng mạnh", "Do hàm băm tự động đổi khi bảng đầy", "Do bảng băm không cho phép thêm phần tử mới", "Do bộ nhớ bị xóa khi bảng đầy"], "A", "PHAN_TICH"),
]

# ===== 10 CÂU LỖI CỐ TÌNH CÀI: (thân, [A,B,C,D], đáp án hoặc None, mã lỗi kỳ vọng) =====
LOI = [
    ("Ngôn ngữ nào dùng để truy vấn cơ sở dữ liệu quan hệ?",
     ["Java", "SQL", "SQL", "HTML"], "B", "TRUNG_PHUONG_AN"),
    ("Cấu trúc dữ liệu nào hoạt động theo nguyên tắc FIFO?",
     ["Stack", "Queue", "queue", "Tree"], "B", "TRUNG_PHUONG_AN"),
    ("Giao thức nào dùng để gửi email?",
     ["SMTP", "FTP", "HTTP", "SMTP"], "A", "TRUNG_PHUONG_AN"),
    ("Python là ngôn ngữ nào sau đây?",
     ["Ngôn ngữ thông dịch", "Ngôn ngữ đa năng", "Có thư viện phong phú", "Tất cả các câu trên đều đúng"], "D", "VI_PHAM_QUY_CHE"),
    ("Đặc điểm nào sau đây của ngôn ngữ Java là đúng?",
     ["Hướng đối tượng", "Chạy trên máy ảo JVM", "Có bộ thu gom rác", "Tất cả các đáp án trên"], "D", "VI_PHAM_QUY_CHE"),
    ("Lệnh nào dùng để xóa một nhánh Git ở máy cục bộ?",
     ["git branch -d", "git add", "git status", "Không có đáp án nào đúng"], "A", "VI_PHAM_QUY_CHE"),
    ("Cổng mặc định của giao thức FTP là bao nhiêu?",
     ["21", "22", "23", "25"], None, "THIEU_DAP_AN"),
    ("Mệnh đề SQL nào dùng để sắp xếp kết quả?",
     ["GROUP BY", "ORDER BY", "HAVING", "WHERE"], None, "THIEU_DAP_AN"),
    ("Thuật toán nào có độ phức tạp trung bình O(n log n)?",
     ["Bubble sort", "Insertion sort", "Merge sort", ""], "C", "THIEU_PHUONG_AN"),
    ("Mô hình nào mô tả 7 tầng trong mạng máy tính?",
     ["TCP/IP", "OSI", "OSI", "ARP"], "B", "TRUNG_PHUONG_AN"),
]

assert len(CHUAN) == 30 and len(LOI) == 10

# Gom thành 40 mục rồi trộn ngẫu nhiên (cố định seed để lần nào cũng giống nhau)
muc = [{"than": t, "pa": pa, "da": da, "bloom": b, "ma_loi": None} for t, pa, da, b in CHUAN]
muc += [{"than": t, "pa": pa, "da": da, "bloom": None, "ma_loi": ml} for t, pa, da, ml in LOI]
random.Random(42).shuffle(muc)

doc = Document()
truth = []
for i, m in enumerate(muc, start=1):
    doc.add_paragraph(f"Câu {i}. {m['than']}")
    for k, noi_dung in zip("ABCD", m["pa"]):
        doc.add_paragraph(f"{k}. {noi_dung}")
    if m["da"]:
        doc.add_paragraph(f"Đáp án: {m['da']}")
    truth.append({"so_cau": i, "co_loi": m["ma_loi"] is not None,
                  "ma_loi": m["ma_loi"], "bloom": m["bloom"]})

doc.save("data/bo_test_40.docx")
with open("data/ground_truth.json", "w", encoding="utf-8") as f:
    json.dump(truth, f, ensure_ascii=False, indent=2)

print("Đã tạo data/bo_test_40.docx và data/ground_truth.json (40 câu)")